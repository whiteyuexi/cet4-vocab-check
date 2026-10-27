#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""用 python-pptx 交叉验证生成的 pptx（与 WPS 渲染互为佐证）
用法：py scripts/pptx-check.py "词汇账户余额查询-答辩.pptx"
"""
import sys
import os
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.expanduser('~'), 'Desktop', '词汇账户余额查询-答辩.pptx')

prs = Presentation(path)
print('文件:', path)
print('尺寸: %.3f × %.3f 英寸 (16:9 = %.4f)' % (
    prs.slide_width / 914400, prs.slide_height / 914400,
    prs.slide_width / prs.slide_height))
print('页数:', len(prs.slides))

total_shapes = total_chars = total_pics = notes_chars = 0
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        total_shapes += 1
        if sh.shape_type is not None and sh.shape_type == 13:
            total_pics += 1
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                total_chars += len(t)
                if not texts:
                    texts.append(t)
    note = ''
    if slide.has_notes_slide:
        note = slide.notes_slide.notes_text_frame.text.strip()
        notes_chars += len(note)
    head = texts[0].replace('\n', ' / ')[:46] if texts else '(无文字)'
    print('  %2d. %-48s 形状 %2d  备注 %3d 字' % (i, head, len(slide.shapes), len(note)))

print('-' * 60)
print('形状总数: %d | 可编辑文字: %d 字 | 图片: %d 张 | 备注合计: %d 字'
      % (total_shapes, total_chars, total_pics, notes_chars))
assert len(prs.slides) == 14, '页数不是 14'
assert total_chars > 2000, '可编辑文字过少，可能退化成图片'
print('结构校验通过：文字是原生可编辑对象，不是图片页。')
