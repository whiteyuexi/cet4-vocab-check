# -*- coding: utf-8 -*-
"""
生成「轻量版」路演稿：把内嵌的演示视频那一页去掉，用于邮件 / 提交门户（有单文件体积上限时）。
视频本身作为独立文件提交，PPT 里保留「视频展示」页说明。

用法: py scripts/make-deck-slim.py <含视频的.pptx> <输出.pptx>
"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

IN = sys.argv[1]
OUT = sys.argv[2]

prs = Presentation(IN)
sldIdLst = prs.slides._sldIdLst
items = list(sldIdLst)

# 找出含 MEDIA 形状的那一页（不按固定页码，避免以后页序变了还删错）
target = None
for i, slide in enumerate(prs.slides):
    for sh in slide.shapes:
        if sh.shape_type is not None and 'MEDIA' in str(sh.shape_type):
            target = i
            break
    if target is not None:
        break

if target is None:
    print('没有找到内嵌媒体页，直接另存一份：' + OUT)
    prs.save(OUT)
    sys.exit(0)

el = items[target]
rId = el.get(qn('r:id'))
prs.part.drop_rel(rId)
sldIdLst.remove(el)
prs.save(OUT)
print('已删除第 %d 页（内嵌视频），剩余 %d 页 → %s' % (target + 1, len(Presentation(OUT).slides), OUT))
