# -*- coding: utf-8 -*-
"""
在用户手改过的《四六级词汇量检测-路演.pptx》里插入一页「未来规划」，
插在「视频展示 / 空白页」之后、「展示完毕，谢谢各位」之前。

为什么用 python-pptx 直接改，而不是重新生成：
  这一版 PPT 是作者自己改过的（学院、署名、第 2 页文案、视频页），
  生成脚本一跑就会把这些改动覆盖掉，所以只能在原文件上做加法。

用法: py scripts/add-roadmap-slide.py [输入.pptx] [输出.pptx]
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

IN = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\27295\Desktop\四六级词汇量检测-路演.pptx"
OUT = sys.argv[2] if len(sys.argv) > 2 else IN

FONT = "微软雅黑"
INK_BG = RGBColor(0x17, 0x17, 0x12)
CARD_BG = RGBColor(0x23, 0x20, 0x1A)
CARD_LN = RGBColor(0x3A, 0x36, 0x2C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0xC9, 0xC2, 0xB4)
ORANGE = RGBColor(0xE8, 0x53, 0x2B)
SALMON = RGBColor(0xFF, 0x8A, 0x6A)


def set_run(run, size, bold=False, color=WHITE, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    # 中文字形要单独指定 a:ea，否则在部分机器上会退回宋体
    rPr = run.font._rPr
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def textbox(slide, x, y, w, h, text, size, bold=False, color=WHITE,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        set_run(p.add_run(), size, bold, color)
        p.runs[0].text = ln
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=None, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.text_frame.word_wrap = True
    return sh


CARDS = [
    ("01", "算法与后端：从「测得准」到「测得细」",
     "错因归类（形近 / 义近 / 搭配）、按信息量调度题目、区间估计升级为 IRT 与贝叶斯；后端支持多设备档案。"),
    ("02", "桌宠：从「会动」到「会教」",
     "按你的答题状态换表情、卡住时主动递一句提示、抓起来喂食换一句「今天先补哪一段」；语音对话当复习搭子。"),
    ("03", "AI：从「解释数据」到「陪你练」",
     "现场出题、批改作文、口语跟读打分；把错因喂回排序，让补词单越用越准，而不是每次都从零开始排。"),
    ("04", "遗忘曲线：从「记录时间」到「预测复测点」",
     "把本地作答和遗忘曲线连起来，预测「哪天该复测哪一段」，而不是只按日历提醒你回来。"),
    ("05", "场景扩展：一个引擎，多种目标",
     "听力词汇、语法覆盖率、专业课词书；教师端班级看板——先用「导出 CSV」顶上，不急着做服务器。"),
    ("06", "形态补全：让它在更多设备上跑",
     "小程序接 AI（云函数转发）、TTS 发音、可选的多设备同步；默认仍然只把作答留在本机。"),
]


def main():
    prs = Presentation(IN)
    before = len(prs.slides)
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    # 背景
    rect(slide, 0, 0, 13.333, 7.5, INK_BG)

    # 标题区
    textbox(slide, 0.85, 0.62, 11.6, 0.3, "未来规划 · 还能往哪走", 11, True, SALMON)
    textbox(slide, 0.85, 0.92, 11.6, 0.78, "这个版本的边界，就是下一版的起点", 28, True, WHITE)
    rect(slide, 0.85, 1.74, 0.52, 0.05, ORANGE)

    # 六块卡片：两列三行
    col_x = [0.85, 6.95]
    row_y = [2.06, 3.42, 4.78]
    cw, ch = 5.53, 1.22
    for idx, (no, title, desc) in enumerate(CARDS):
        x = col_x[idx % 2]
        y = row_y[idx // 2]
        rect(slide, x, y, cw, ch, CARD_BG, CARD_LN, 0.06, MSO_SHAPE.ROUNDED_RECTANGLE)
        textbox(slide, x + 0.22, y + 0.16, 0.42, 0.3, no, 11, True, ORANGE)
        textbox(slide, x + 0.66, y + 0.13, cw - 0.9, 0.34, title, 13, True, WHITE)
        textbox(slide, x + 0.22, y + 0.55, cw - 0.46, 0.6, desc, 10, False, DIM, spacing=1.25)

    # 收尾一句：给全场一个落点
    textbox(slide, 0.85, 6.28, 0.28, 0.34, "—", 15, True, ORANGE)
    textbox(slide, 1.16, 6.24, 11.3, 0.62,
            "下一版的目标很具体：让它在你想放弃的时候说一句「再答 10 个词，就到中点了」。",
            13.5, True, WHITE)

    # 演讲备注
    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "收尾这页不要念完六块，挑两块讲透：① 桌宠从「会动」到「会教」——现在它只提供情绪价值，"
        "下一步让它按你的答题状态说话、卡住时递提示；② 算法与后端——错因归类 + 更细的题目调度，"
        "让补词单越用越准。最后停在底下一句话上："
        "让它在你想放弃的时候说一句「再答 10 个词就到中点」，这就是这个工具想做到的事。"
    )

    # 移到末尾前一页（放在「展示完毕，谢谢各位」之前）
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    new_id = ids[-1]
    sldIdLst.remove(new_id)
    sldIdLst.insert(len(ids) - 2, new_id)

    prs.save(OUT)
    print("已插入 1 页，共 %d → %d 页：%s" % (before, len(Presentation(OUT).slides), OUT))


if __name__ == "__main__":
    main()
